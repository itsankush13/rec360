"""
PPTX as a CV format.

A CV built in PowerPoint puts its real content in three places: shape text
frames (name, headline), tables (skills grids, experience rows) and slide
notes (a recruiter's private annotation, or a copy the candidate hid there).
Missing any one of the three yields a near-empty document that then fails
the MIN_TEXT_CHARS gate for the wrong reason.

Fixtures build real PPTX bytes with python-pptx rather than mocking the
parser, matching the DOCX and PDF fixtures in test_processing.py and
test_document_quality.py — the point of this validation layer is what
actually happens when python-pptx meets a real or a broken file.
"""
import pytest

from app.core import document_intake
from app.core.document_intake import DocumentRejected


def make_pptx(slides_text=None, table_rows=None, notes=None):
    """
    Build a .pptx in memory.

    slides_text: list of strings, one text-box slide per string.
    table_rows: rows (each a list of cell strings) added to a table on a
        new slide, appended after the text slides.
    notes: string, put in the notes of the first slide.
    """
    import pptx
    from pptx.util import Inches

    presentation = pptx.Presentation()
    blank_layout = presentation.slide_layouts[6]

    slides = []
    for text in slides_text or []:
        slide = presentation.slides.add_slide(blank_layout)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        box.text_frame.text = text
        slides.append(slide)

    if table_rows:
        slide = presentation.slides.add_slide(blank_layout)
        rows, cols = len(table_rows), len(table_rows[0])
        graphic_frame = slide.shapes.add_table(
            rows, cols, Inches(1), Inches(1), Inches(6), Inches(2)
        )
        table = graphic_frame.table
        for r, row_values in enumerate(table_rows):
            for c, value in enumerate(row_values):
                table.cell(r, c).text = value
        slides.append(slide)

    if notes and slides:
        slides[0].notes_slide.notes_text_frame.text = notes

    import io
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------

def test_pptx_with_text_in_shapes_extracts(tmp_path):
    data = make_pptx(slides_text=[
        "Priya Menon\npriya.menon@example.com\n+91 98765 43210",
        "EXPERIENCE\nStaff Engineer, Acme Corp (2020-present). "
        "Led the payments platform team across three continents for years.",
    ])
    path = tmp_path / "cv.pptx"
    path.write_bytes(data)

    extracted = document_intake.extract_document(path, "cv.pptx")

    assert "Priya Menon" in extracted.text
    assert "priya.menon@example.com" in extracted.text
    assert "Staff Engineer" in extracted.text


def test_pptx_with_text_in_tables_extracts(tmp_path):
    data = make_pptx(
        slides_text=[
            "Priya Menon\npriya.menon@example.com\n+91 98765 43210\n"
            "Senior backend engineer with eight years building Python "
            "services on AWS across multiple large scale platforms.",
        ],
        table_rows=[
            ["Skill", "Years"],
            ["Python", "8"],
            ["PostgreSQL", "6"],
        ],
    )
    path = tmp_path / "cv.pptx"
    path.write_bytes(data)

    extracted = document_intake.extract_document(path, "cv.pptx")

    assert "PostgreSQL" in extracted.text
    assert "Python" in extracted.text


def test_pptx_slide_notes_extract(tmp_path):
    data = make_pptx(
        slides_text=[
            "Priya Menon\npriya.menon@example.com\n+91 98765 43210\n"
            "Senior backend engineer with eight years building Python "
            "services on AWS across multiple large scale platforms.",
        ],
        notes="Referred by Arun Verma, prior colleague at Globex.",
    )
    path = tmp_path / "cv.pptx"
    path.write_bytes(data)

    extracted = document_intake.extract_document(path, "cv.pptx")

    assert "Referred by Arun Verma" in extracted.text


def test_pptx_slide_to_page_mapping_is_one_based(tmp_path):
    data = make_pptx(slides_text=[
        "Priya Menon\npriya.menon@example.com\n+91 98765 43210\n"
        "Senior backend engineer with eight years building Python "
        "services on AWS across multiple large scale platforms.",
        "EDUCATION\nB.Tech Computer Science, NIT Trichy, graduated with "
        "distinction and multiple academic awards over four years.",
    ])
    path = tmp_path / "cv.pptx"
    path.write_bytes(data)

    extracted = document_intake.extract_document(path, "cv.pptx")

    assert extracted.page_count == 2
    assert len(extracted.pages) == 2
    assert "Priya Menon" in extracted.pages[0]
    assert "EDUCATION" in extracted.pages[1]


def test_corrupt_pptx_raises_document_rejected(tmp_path):
    path = tmp_path / "broken.pptx"
    path.write_bytes(b"definitely not a zip archive")

    with pytest.raises(DocumentRejected) as excinfo:
        document_intake.extract_document(path, "broken.pptx")

    assert excinfo.value.code == document_intake.JobErrorCode.CORRUPT_FILE


def test_ppt_extension_is_still_rejected_naming_pptx(tmp_path):
    path = tmp_path / "old.ppt"
    path.write_bytes(b"not a real ppt file")

    with pytest.raises(DocumentRejected) as excinfo:
        document_intake.extract_document(path, "old.ppt")

    assert excinfo.value.code == document_intake.JobErrorCode.UNSUPPORTED_FORMAT
    assert ".pptx" in excinfo.value.message


def test_ppt_is_still_rejected_on_upload_validation():
    with pytest.raises(DocumentRejected) as excinfo:
        document_intake.validate_upload("old.ppt", b"not a real ppt file")

    assert excinfo.value.code == document_intake.JobErrorCode.UNSUPPORTED_FORMAT
    assert ".pptx" in excinfo.value.message
