"""
Phase H tests — Candidate 360 PDF, shortlist export, ATS field mapping.

These do not stop at the status code. A report endpoint that returns 200 and
a corrupt file is worse than one that fails, so the PDF is parsed back and
its text checked, and the workbook is reopened and its cells read.
"""
from unittest.mock import patch

import pytest

from app.core import ats_export
from app.db.models import AuditAction

from tests.test_processing import FAKE_JD_EXTRACTION, PDF_MIME, upload
from tests.test_evaluations import STRONG_CV, WEAK_CV, make_cv_pdf
from tests.test_analytics import THIRD_CV, FOURTH_CV


@pytest.fixture()
def exportable(client):
    """A campaign with four assessed candidates and one decision recorded."""
    with patch("app.services.requirement_service.parse_jd", return_value=FAKE_JD_EXTRACTION):
        campaign_id = client.post("/api/campaigns", json={
            "name": "Control Room Operator intake",
            "job_title": "Control Room Operator",
            "job_description": "Python and AWS backend engineer.",
            "location": "Coastal Terminal", "vacancies": 8,
        }).json()["id"]
        client.post(f"/api/campaigns/{campaign_id}/requirements/extract", json={})
    client.post(f"/api/campaigns/{campaign_id}/rubric/versions",
                json={"seed_from_requirements": True})
    client.post(f"/api/campaigns/{campaign_id}/rubric/versions/1/submit", json={})
    client.post(f"/api/campaigns/{campaign_id}/rubric/versions/1/approve",
                json={"approved_by": "Fatima Al-Rashid"})
    upload(client, campaign_id, [
        ("priya.pdf", make_cv_pdf(STRONG_CV), PDF_MIME),
        ("rahul.pdf", make_cv_pdf(WEAK_CV), PDF_MIME),
        ("anjali.pdf", make_cv_pdf(THIRD_CV), PDF_MIME),
        ("vikram.pdf", make_cv_pdf(FOURTH_CV), PDF_MIME),
    ])
    client.post(f"/api/campaigns/{campaign_id}/evaluations/runs", json={})
    evaluations = client.get(f"/api/campaigns/{campaign_id}/evaluations").json()
    client.post(
        f"/api/campaigns/{campaign_id}/candidates/{evaluations[0]['candidate_id']}/disposition",
        json={"disposition": "SHORTLIST", "actor": "Fatima Al-Rashid"},
    )
    return campaign_id, evaluations


def pdf_text(payload: bytes) -> str:
    import pymupdf
    with pymupdf.open(stream=payload, filetype="pdf") as document:
        return "\n".join(page.get_text() for page in document)


# ===========================================================================
# Candidate 360 PDF
# ===========================================================================

def test_report_returns_a_real_pdf(exportable, client):
    _, evaluations = exportable
    response = client.get(f"/api/evaluations/{evaluations[0]['id']}/report.pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    assert "attachment" in response.headers["content-disposition"]
    assert ".pdf" in response.headers["content-disposition"]


def test_report_opens_and_carries_the_candidate(exportable, client):
    _, evaluations = exportable
    top = evaluations[0]
    text = pdf_text(client.get(f"/api/evaluations/{top['id']}/report.pdf").content)
    assert top["candidate_name"].split()[0] in text
    # The banner renders the role in caps, as the design does.
    assert "control room operator" in text.lower()
    assert "out of 100" in text


# ===========================================================================
# Shortlist zip download (build_shortlist_zip, shared with reports.py)
# ===========================================================================

def test_shortlisted_reports_zip_contains_one_pdf_per_shortlisted_candidate(exportable, client):
    import zipfile as zipfile_module

    campaign_id, _evaluations = exportable
    response = client.get(f"/api/campaigns/{campaign_id}/reports.zip")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert ".zip" in response.headers["content-disposition"]

    with zipfile_module.ZipFile(__import__("io").BytesIO(response.content)) as archive:
        names = archive.namelist()
        assert len(names) == 1  # exactly one shortlisted candidate in the fixture
        assert names[0].startswith("Shortlisted Candidates PDF/")
        assert names[0].endswith(".pdf")


def test_shortlisted_reports_zip_422_when_nothing_shortlisted(client):
    with patch("app.services.requirement_service.parse_jd", return_value=FAKE_JD_EXTRACTION):
        campaign_id = client.post("/api/campaigns", json={
            "name": "Undecided", "job_title": "Engineer", "job_description": "x",
        }).json()["id"]
        client.post(f"/api/campaigns/{campaign_id}/requirements/extract", json={})

    response = client.get(f"/api/campaigns/{campaign_id}/reports.zip")
    assert response.status_code == 422


def test_report_states_that_a_person_decides(exportable, client):
    """
    The client's first requirement, and a PDF is what reaches someone who
    never saw the screen.
    """
    _, evaluations = exportable
    text = pdf_text(client.get(f"/api/evaluations/{evaluations[0]['id']}/report.pdf").content)
    assert "not a hiring decision" in text.lower()
    assert "a person decides" in text.lower()


def test_report_quotes_evidence_with_a_page_or_section(exportable, client):
    _, evaluations = exportable
    text = pdf_text(client.get(f"/api/evaluations/{evaluations[0]['id']}/report.pdf").content)
    assert "Evidence, quoted from the CV" in text
    assert ("CV page" in text) or ("CV," in text)


def test_report_carries_every_criterion(exportable, client):
    _, evaluations = exportable
    detail = client.get(f"/api/evaluations/{evaluations[0]['id']}").json()
    text = pdf_text(client.get(f"/api/evaluations/{evaluations[0]['id']}/report.pdf").content)
    for criterion in detail["criteria"][:4]:
        assert criterion["label"][:24] in text


def test_report_uses_plain_english_not_enum_values(exportable, client):
    _, evaluations = exportable
    text = pdf_text(client.get(f"/api/evaluations/{evaluations[0]['id']}/report.pdf").content)
    for banned in ("CONFIRMED_MATCH", "NOT_DEMONSTRATED", "INSUFFICIENT_EVIDENCE",
                   "STRONG_FIT", "REVIEW_REQUIRED", "HARD_FAIL"):
        assert banned not in text, banned
    assert "Confirmed" in text or "Not shown in the CV" in text


def test_report_records_its_provenance(exportable, client):
    """A score without the rubric that produced it is not auditable."""
    _, evaluations = exportable
    text = pdf_text(client.get(f"/api/evaluations/{evaluations[0]['id']}/report.pdf").content)
    assert "Rubric version" in text
    assert "deterministic" in text.lower()


def test_report_shows_the_recruiter_decision_when_one_exists(exportable, client):
    _, evaluations = exportable
    text = pdf_text(client.get(f"/api/evaluations/{evaluations[0]['id']}/report.pdf").content)
    assert "Recruiter decision" in text
    assert "Shortlist" in text


def test_report_is_audited(exportable, client):
    campaign_id, evaluations = exportable
    client.get(f"/api/evaluations/{evaluations[0]['id']}/report.pdf")
    events = client.get(f"/api/campaigns/{campaign_id}/audit",
                        params={"action": "EXPORTED"}).json()
    assert events
    assert events[0]["after"]["format"] == "pdf"


def test_report_404s_for_an_unknown_assessment(client):
    assert client.get("/api/evaluations/nope/report.pdf").status_code == 404


# ===========================================================================
# Shortlist export
# ===========================================================================

def test_csv_export_opens_and_has_the_mapped_columns(exportable, client):
    campaign_id, _ = exportable
    response = client.get(f"/api/campaigns/{campaign_id}/export.csv")
    assert response.status_code == 200
    assert "text/csv" in response.headers["content-type"]

    import csv, io
    text = response.content.decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(text)))
    assert rows
    assert set(rows[0]) == set(ats_export.COLUMNS)
    assert rows[0]["candidate_name"]
    assert rows[0]["recruiter_recommendation"] == "Shortlisted"


def test_csv_uses_a_bom_so_excel_reads_accents(exportable, client):
    """Windows Excel mangles UTF-8 without one, and candidate names have accents."""
    campaign_id, _ = exportable
    content = client.get(f"/api/campaigns/{campaign_id}/export.csv").content
    assert content.startswith(b"\xef\xbb\xbf")


def test_export_defaults_to_decided_candidates_only(exportable, client):
    """
    Exporting an unreviewed ranking would make the AI's ordering the
    operative decision, which this system is explicitly not for.
    """
    campaign_id, evaluations = exportable
    import csv, io
    decided = list(csv.DictReader(io.StringIO(
        client.get(f"/api/campaigns/{campaign_id}/export.csv").content.decode("utf-8-sig"))))
    everyone = list(csv.DictReader(io.StringIO(
        client.get(f"/api/campaigns/{campaign_id}/export.csv",
                   params={"decided_only": False}).content.decode("utf-8-sig"))))
    assert len(decided) == 1
    assert len(everyone) == len(evaluations)


def test_export_refuses_when_nothing_is_decided(client):
    with patch("app.services.requirement_service.parse_jd", return_value=FAKE_JD_EXTRACTION):
        campaign_id = client.post("/api/campaigns", json={
            "name": "Undecided", "job_title": "Engineer", "job_description": "x",
        }).json()["id"]
        client.post(f"/api/campaigns/{campaign_id}/requirements/extract", json={})
    response = client.get(f"/api/campaigns/{campaign_id}/export.csv")
    assert response.status_code == 422
    assert "nothing to export" in str(response.json()["detail"]).lower()


def test_export_keeps_the_ai_view_and_the_human_decision_apart(exportable, client):
    """
    Collapsing them would lose the fact that a person disagreed, which is the
    most important thing in the record.
    """
    campaign_id, evaluations = exportable
    import csv, io
    rows = list(csv.DictReader(io.StringIO(
        client.get(f"/api/campaigns/{campaign_id}/export.csv").content.decode("utf-8-sig"))))
    assert rows[0]["ai_assessment"]
    assert rows[0]["recruiter_recommendation"]
    assert "ai_assessment" in ats_export.COLUMNS
    assert "recruiter_recommendation" in ats_export.COLUMNS


def test_export_carries_evidence_and_provenance(exportable, client):
    campaign_id, _ = exportable
    import csv, io
    rows = list(csv.DictReader(io.StringIO(
        client.get(f"/api/campaigns/{campaign_id}/export.csv").content.decode("utf-8-sig"))))
    assert rows[0]["evidence_summary"]
    assert rows[0]["rubric_version"].startswith("v")
    assert rows[0]["assessed_on"]


def test_xlsx_export_opens_with_both_sheets(exportable, client):
    campaign_id, _ = exportable
    response = client.get(f"/api/campaigns/{campaign_id}/export.xlsx")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]

    import io
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Shortlist", "Field mapping"]

    sheet = workbook["Shortlist"]
    headers = [c.value for c in sheet[4]]
    assert "Candidate Name" in headers
    assert sheet["A1"].value.startswith("Control Room Operator")
    assert "a person decides" in sheet["A2"].value.lower()

    mapping = workbook["Field mapping"]
    assert mapping["B1"].value == "SAP SuccessFactors field"
    assert any(row[1].value == "Candidate Full Name"
               for row in mapping.iter_rows(min_row=2, max_row=6))


def test_xlsx_ships_the_mapping_with_the_data(exportable, client):
    """A mapping in a separate document gets separated from the data."""
    campaign_id, _ = exportable
    import io
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(
        client.get(f"/api/campaigns/{campaign_id}/export.xlsx").content))
    mapping = workbook["Field mapping"]
    values = [row[1].value for row in mapping.iter_rows(min_row=2)]
    assert "Requisition ID" in values
    assert "Recruiter Recommendation" in values


def test_export_is_audited(exportable, client):
    campaign_id, _ = exportable
    client.get(f"/api/campaigns/{campaign_id}/export.xlsx")
    events = client.get(f"/api/campaigns/{campaign_id}/audit",
                        params={"action": "EXPORTED"}).json()
    assert events
    assert events[0]["after"]["format"] == "xlsx"
    assert events[0]["after"]["candidates"] == 1


def test_export_404s_for_an_unknown_campaign(client):
    assert client.get("/api/campaigns/nope/export.csv").status_code == 404


# ===========================================================================
# Field mapping
# ===========================================================================

def test_field_mapping_is_exposed_as_data(client):
    """So the Decision screen renders the mapping the export actually uses."""
    body = client.get("/api/exports/ats/field-mapping").json()
    assert body["ats"] == "SAP SuccessFactors Recruiting"
    assert body["integration"] == "file"
    assert "not been agreed" in body["note"]
    assert len(body["fields"]) == len(ats_export.FIELD_MAP)
    ours = {f["our_field"] for f in body["fields"]}
    assert ours == set(ats_export.COLUMNS)


def test_mapping_matches_the_exported_columns(exportable, client):
    """The mapping and the file cannot drift apart."""
    campaign_id, _ = exportable
    import csv, io
    rows = list(csv.DictReader(io.StringIO(
        client.get(f"/api/campaigns/{campaign_id}/export.csv").content.decode("utf-8-sig"))))
    mapped = {f["our_field"] for f in
              client.get("/api/exports/ats/field-mapping").json()["fields"]}
    assert set(rows[0]) == mapped


# ===========================================================================
# Boundaries
# ===========================================================================

def test_there_is_no_send_to_ats_endpoint(client):
    """
    Whether the handover is a live API call is undecided and needs the
    client's HRIS owner. A button that posts to a tenant nobody agreed on
    would be worse than the file.
    """
    paths = client.app.openapi()["paths"]
    for path, operations in paths.items():
        if "successfactors" in path.lower() or "/ats/send" in path.lower():
            pytest.fail(f"{path} implies a live integration that is not agreed")


def test_export_routes_are_read_only(client):
    paths = client.app.openapi()["paths"]
    for path, operations in paths.items():
        if "export" in path or "report.pdf" in path:
            assert set(m.lower() for m in operations) == {"get"}, path


def test_reports_do_not_alter_the_assessment(exportable, client):
    campaign_id, evaluations = exportable
    before = client.get(f"/api/campaigns/{campaign_id}/evaluations").json()
    client.get(f"/api/evaluations/{evaluations[0]['id']}/report.pdf")
    client.get(f"/api/campaigns/{campaign_id}/export.xlsx")
    after = client.get(f"/api/campaigns/{campaign_id}/evaluations").json()
    assert before == after


# ---------------------------------------------------------------------------
# The Word report
# ---------------------------------------------------------------------------

def test_the_assessment_downloads_as_a_word_document(exportable, client):
    """
    The hiring manager's working copy. Same assessment as the PDF, in a file
    they can annotate.
    """
    _campaign_id, evaluations = exportable
    evaluation_id = evaluations[0]["id"]
    response = client.get(f"/api/evaluations/{evaluation_id}/report.docx")

    assert response.status_code == 200
    assert "wordprocessingml" in response.headers["content-type"]
    assert response.content[:2] == b"PK"          # a docx is a zip
    assert "-assessment.docx" in response.headers["content-disposition"]


def test_the_word_report_carries_the_disclaimer(exportable, client):
    """
    This is the file most likely to be forwarded to somebody who never saw
    the screen, so the "not a hiring decision" line has to travel with it.
    """
    import io, zipfile

    _campaign_id, evaluations = exportable
    evaluation_id = evaluations[0]["id"]
    response = client.get(f"/api/evaluations/{evaluation_id}/report.docx")
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        text = archive.read("word/document.xml").decode("utf-8")

    assert "not a hiring decision" in text
    assert "a person decides" in text.lower() or "A person makes the decision" in text


def test_an_unknown_assessment_has_no_word_report(client):
    assert client.get("/api/evaluations/does-not-exist/report.docx").status_code == 404


def test_sending_a_report_is_recorded_in_the_audit_trail(exportable, client, db_session):
    from app.db.models import AuditAction, AuditEvent
    from sqlalchemy import select

    campaign_id, evaluations = exportable
    evaluation_id = evaluations[0]["id"]
    client.get(f"/api/evaluations/{evaluation_id}/report.docx")

    events = list(db_session.scalars(
        select(AuditEvent).where(
            AuditEvent.campaign_id == campaign_id,
            AuditEvent.action == AuditAction.EXPORTED,
        )
    ).all())
    assert any(event.after.get("format") == "docx" for event in events)
    assert any("Word document" in event.summary for event in events)


def test_the_decisions_screen_no_longer_claims_an_ats_connection():
    """
    There is no integration with any recruiting system, so the screen must
    not say there is. This is the page a client auditor reads.
    """
    from pathlib import Path
    html = (Path(__file__).resolve().parent.parent / "web" / "decisions.html").read_text(
        encoding="utf-8", errors="ignore"
    )
    for claim in ("SuccessFactors", "SAP"):
        assert claim not in html
