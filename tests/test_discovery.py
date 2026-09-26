"""
`/discovery/resolve` and `/discovery/import` — paste a folder path, see the
CVs there, import the ones picked.

`/discovery/import` must reuse the same intake path as
`POST /api/campaigns/{id}/batches` (see tests/test_processing.py), so
duplicate detection, exception codes and audit rows read identically
whichever door a CV came through.
"""
from unittest.mock import patch

import pytest

FAKE_JD_EXTRACTION = {
    "role_title": "Senior Software Engineer",
    "required_skills": ["Python", "AWS", "PostgreSQL"],
    "preferred_skills": ["Kubernetes"],
    "min_experience_years": 5,
    "education_requirement": "Bachelor's in Computer Science",
    "key_responsibilities": ["Own backend services"],
    "seniority_level": "Senior",
}


@pytest.fixture()
def screening_campaign(client):
    with patch("app.services.requirement_service.parse_jd", return_value=FAKE_JD_EXTRACTION):
        campaign_id = client.post("/api/campaigns", json={
            "name": "Discovery campaign",
            "job_title": "Senior Software Engineer",
            "job_description": "Python and AWS backend engineer.",
        }).json()["id"]
        client.post(f"/api/campaigns/{campaign_id}/requirements/extract", json={})

    client.post(
        f"/api/campaigns/{campaign_id}/rubric/versions",
        json={"seed_from_requirements": True},
    )
    client.post(f"/api/campaigns/{campaign_id}/rubric/versions/1/submit", json={})
    approved = client.post(
        f"/api/campaigns/{campaign_id}/rubric/versions/1/approve", json={"approved_by": "Priya Nair"}
    )
    assert approved.status_code == 200, approved.text
    return campaign_id


def test_resolve_lists_files_in_a_local_folder(client, tmp_path):
    (tmp_path / "candidate.pdf").write_bytes(b"%PDF-1.4 fake")
    (tmp_path / "notes.txt").write_text("not a cv format")

    response = client.post("/discovery/resolve", json={"location": str(tmp_path)})

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["resolved_folder"] == str(tmp_path)
    assert [f["filename"] for f in body["files"]] == ["candidate.pdf"]


def test_resolve_unresolvable_link_names_searched_roots(client, tmp_path, monkeypatch):
    import app.api.discovery as discovery
    from app.core.cv_source import SyncedFolderSource

    fake_root = tmp_path / "OneDrive - Nobody"
    fake_root.mkdir()
    monkeypatch.setattr(discovery, "_source", SyncedFolderSource(sync_roots=[fake_root]))

    response = client.post("/discovery/resolve", json={
        "location": "https://nowhere.sharepoint.com/sites/Missing/Shared%20Documents",
    })

    assert response.status_code == 422
    detail = response.json()["detail"]
    assert str(fake_root) in detail["searched_roots"]


def test_import_rejects_a_path_outside_the_resolved_folder(client, screening_campaign, tmp_path):
    resolved_root = tmp_path / "resolved"
    resolved_root.mkdir()
    outside_dir = tmp_path / "outside"
    outside_dir.mkdir()
    outsider = outside_dir / "sneaky.pdf"
    outsider.write_bytes(b"%PDF-1.4 fake")

    response = client.post("/discovery/import", json={
        "campaign_id": screening_campaign,
        "resolved_root": str(resolved_root),
        "paths": [str(outsider)],
    })

    assert response.status_code == 400


def test_resolve_without_campaign_id_does_not_rank(client, tmp_path):
    (tmp_path / "candidate.pdf").write_bytes(b"%PDF-1.4 fake")

    response = client.post("/discovery/resolve", json={"location": str(tmp_path)})

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["ranking_available"] is False
    assert body["files"][0]["relevance_score"] is None
    assert body["files"][0]["recommended"] is False


def test_resolve_ranks_by_rubric_term_coverage(client, screening_campaign, tmp_path):
    import docx

    def _write(name: str, lines: list[str]) -> None:
        document = docx.Document()
        for line in lines:
            document.add_paragraph(line)
        document.save(str(tmp_path / name))

    _write("strong-match.docx", [
        "Jordan Lee", "jordan.lee@example.com", "SUMMARY",
        "Senior backend engineer building Python services on AWS with PostgreSQL "
        "and Kubernetes for container orchestration across eight years.",
    ])
    _write("no-match.docx", [
        "Sam Rivera", "sam.rivera@example.com", "SUMMARY",
        "Marketing coordinator running social media campaigns, email newsletters "
        "and in-store promotions for a regional retail chain over the past six years.",
    ])

    response = client.post("/discovery/resolve", json={
        "location": str(tmp_path),
        "campaign_id": screening_campaign,
    })

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["ranking_available"] is True

    by_name = {f["filename"]: f for f in body["files"]}
    strong = by_name["strong-match.docx"]
    weak = by_name["no-match.docx"]

    assert strong["relevance_score"] > weak["relevance_score"]
    assert strong["recommended"] is True
    assert weak["recommended"] is False
    # Ranked descending: the strong match leads the list.
    assert body["files"][0]["filename"] == "strong-match.docx"


def test_resolve_ranking_unavailable_without_approved_rubric(client, tmp_path):
    campaign_id = client.post("/api/campaigns", json={
        "name": "No rubric yet",
        "job_title": "Engineer",
        "job_description": "Placeholder.",
    }).json()["id"]
    (tmp_path / "candidate.pdf").write_bytes(b"%PDF-1.4 fake")

    response = client.post("/discovery/resolve", json={
        "location": str(tmp_path),
        "campaign_id": campaign_id,
    })

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["ranking_available"] is False
    assert body["files"][0]["relevance_score"] is None


def test_import_reuses_the_batch_intake_path(client, screening_campaign, tmp_path):
    import docx

    resolved_root = tmp_path / "resolved"
    resolved_root.mkdir()
    document = docx.Document()
    document.add_paragraph("Priya Menon")
    document.add_paragraph("priya.menon@example.com")
    document.add_paragraph("+91 98765 43210")
    document.add_paragraph("SUMMARY")
    document.add_paragraph("Senior backend engineer with eight years building Python services.")
    cv_path = resolved_root / "priya.docx"
    document.save(str(cv_path))

    response = client.post("/discovery/import", json={
        "campaign_id": screening_campaign,
        "resolved_root": str(resolved_root),
        "paths": [str(cv_path)],
    })

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["accepted_files"] == 1
    assert body["batch"]["campaign_id"] == screening_campaign
